import os
import sys
import json
import re
import socket
import threading
import time
import traceback
from datetime import datetime

import customtkinter as ctk

from tkinter import (
    filedialog,
    messagebox
)

from gst_client import GSTClient

from excel_handler import (
    read_gstins,
    save_taxpayer_records,
    save_return_records
)


# ============================================================
# APPLICATION SETTINGS
# ============================================================

APP_NAME = "GST Bulk Search - by Rahul Pal"
APP_VERSION = "1.0"
CONTACT_NUMBER = "+91 7888840836"
GST_HOST = "services.gst.gov.in"


def resource_path(relative_path):
    """
    Returns the correct path for normal Python execution
    and for a PyInstaller packaged application.
    """
    if hasattr(sys, "_MEIPASS"):
        base_path = sys._MEIPASS
    else:
        base_path = os.path.dirname(
            os.path.abspath(__file__)
        )

    return os.path.join(
        base_path,
        relative_path
    )


def application_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


SETTINGS_PATH = os.path.join(application_dir(), "settings.json")
DEFAULT_SETTINGS = {
    "request_delay": 1.0,
    "retry_count": 4,
    "output_folder": os.path.join(application_dir(), "output"),
    "auto_open_excel": False,
    "theme": "System",
}


def load_settings():
    settings = DEFAULT_SETTINGS.copy()
    try:
        if os.path.exists(SETTINGS_PATH):
            with open(SETTINGS_PATH, "r", encoding="utf-8") as f:
                saved = json.load(f)
            if isinstance(saved, dict):
                settings.update(saved)
    except Exception:
        pass
    return settings


def is_valid_gstin(value):
    value = str(value or "").strip().upper()
    return bool(re.fullmatch(
        r"[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][0-9A-Z]Z[0-9A-Z]", value
    ))


def validate_gstins(values):
    valid, invalid, duplicates = [], [], []
    seen = set()
    for raw in values:
        value = str(raw or "").strip().upper()
        if not value:
            continue
        if not is_valid_gstin(value):
            invalid.append(value)
        elif value in seen:
            duplicates.append(value)
        else:
            seen.add(value)
            valid.append(value)
    return valid, invalid, duplicates


MODEL_PATH = resource_path(
    "captcha_model.onnx"
)

REQUEST_DELAY = 1.0


MONTHS = [
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
    "January",
    "February",
    "March"
]


def build_financial_years():
    current_year = datetime.now().year

    years = []

    for year in range(
        current_year + 1,
        2017,
        -1
    ):
        years.append(
            f"{year}-{str(year + 1)[-2:]}"
        )

    return years


FINANCIAL_YEARS = build_financial_years()


# ============================================================
# MAIN APPLICATION
# ============================================================

class GSTBulkApp(ctk.CTk):

    def __init__(self):
        super().__init__()

        self.settings = load_settings()
        ctk.set_appearance_mode(self.settings.get("theme", "System"))

        self.title(APP_NAME)

        self.geometry(
            "900x700"
        )

        self.minsize(
            850,
            650
        )

        ctk.set_default_color_theme(
            "blue"
        )

        self.input_file = ""

        self.processing = False
        self.stop_requested = False
        self.last_output_path = None

        self.client = None
        self.last_error_details = ""
        self.taxpayer_valid_gstins = None
        self.return_valid_gstins = None

        self.create_ui()
        self.after(600, self._start_network_check)

    # ========================================================
    # UI
    # ========================================================

    def create_ui(self):

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        title = ctk.CTkLabel(
            self,
            text="GST BULK SEARCH",
            font=ctk.CTkFont(
                size=28,
                weight="bold"
            )
        )

        title.pack(
            pady=(25, 0)
        )

        subtitle = ctk.CTkLabel(
            self,
            text="Developed by Rahul Pal",
            font=ctk.CTkFont(
                size=14
            )
        )

        subtitle.pack(
            pady=(2, 20)
        )

        header_actions = ctk.CTkFrame(self, fg_color="transparent")
        header_actions.pack(fill="x", padx=30, pady=(0, 5))

        self.network_label = ctk.CTkLabel(
            header_actions, text="● Checking network...",
            font=ctk.CTkFont(size=13, weight="bold")
        )
        self.network_label.pack(side="left")

        ctk.CTkButton(header_actions, text="How to Use", width=105,
                      command=self.show_how_to).pack(side="right", padx=(6, 0))
        ctk.CTkButton(header_actions, text="About", width=80,
                      command=self.show_about).pack(side="right", padx=(6, 0))
        ctk.CTkButton(header_actions, text="Settings", width=90,
                      command=self.show_settings).pack(side="right", padx=(6, 0))

        # ----------------------------------------------------
        # TABS
        # ----------------------------------------------------

        self.tabs = ctk.CTkTabview(
            self
        )

        self.tabs.pack(
            fill="both",
            expand=True,
            padx=25,
            pady=10
        )

        self.taxpayer_tab = (
            self.tabs.add(
                "Bulk Taxpayer Details"
            )
        )

        self.return_tab = (
            self.tabs.add(
                "Return Filing Status"
            )
        )

        self.build_taxpayer_tab()
        self.build_return_tab()

        # ----------------------------------------------------
        # LOG
        # ----------------------------------------------------

        log_label = ctk.CTkLabel(
            self,
            text="Processing Log",
            anchor="w"
        )

        log_label.pack(
            fill="x",
            padx=30
        )

        self.log_box = ctk.CTkTextbox(
            self,
            height=130
        )

        self.log_box.pack(
            fill="x",
            padx=30,
            pady=(5, 15)
        )

        self.log_box.configure(
            state="disabled"
        )

    # ========================================================
    # TAXPAYER TAB
    # ========================================================

    def build_taxpayer_tab(self):

        frame = ctk.CTkFrame(
            self.taxpayer_tab
        )

        frame.pack(
            fill="both",
            expand=True,
            padx=20,
            pady=20
        )

        label = ctk.CTkLabel(
            frame,
            text=(
                "Select an Excel file containing "
                "a GSTIN column."
            ),
            font=ctk.CTkFont(
                size=16
            )
        )

        label.pack(
            pady=(25, 15)
        )

        self.taxpayer_file_label = (
            ctk.CTkLabel(
                frame,
                text="No file selected"
            )
        )

        self.taxpayer_file_label.pack(
            pady=5
        )

        browse = ctk.CTkButton(
            frame,
            text="Select Excel File",
            command=self.select_taxpayer_file
        )

        browse.pack(
            pady=10
        )

        self.taxpayer_validation_label = ctk.CTkLabel(
            frame, text="Select a file to validate GSTINs."
        )
        self.taxpayer_validation_label.pack(pady=(0, 4))

        self.taxpayer_progress = (
            ctk.CTkProgressBar(
                frame,
                width=500
            )
        )

        self.taxpayer_progress.set(
            0
        )

        self.taxpayer_progress.pack(
            pady=(30, 5)
        )

        self.taxpayer_progress_label = (
            ctk.CTkLabel(
                frame,
                text="0 / 0"
            )
        )

        self.taxpayer_progress_label.pack()

        self.taxpayer_start_button = (
            ctk.CTkButton(
                frame,
                text="START PROCESSING",
                width=220,
                height=42,
                command=self.start_taxpayer
            )
        )

        self.taxpayer_start_button.pack(
            pady=25
        )

        taxpayer_actions = ctk.CTkFrame(frame, fg_color="transparent")
        taxpayer_actions.pack(pady=(0, 15))

        self.taxpayer_stop_button = ctk.CTkButton(
            taxpayer_actions, text="STOP / CANCEL", width=160,
            fg_color="#B3261E", hover_color="#8C1D18",
            state="disabled", command=self.request_stop
        )
        self.taxpayer_stop_button.grid(row=0, column=0, padx=6)

        self.taxpayer_open_file_button = ctk.CTkButton(
            taxpayer_actions, text="Open Output File", width=160,
            state="disabled", command=self.open_output_file
        )
        self.taxpayer_open_file_button.grid(row=0, column=1, padx=6)

        self.taxpayer_open_folder_button = ctk.CTkButton(
            taxpayer_actions, text="Open Output Folder", width=170,
            state="disabled", command=self.open_output_folder
        )
        self.taxpayer_open_folder_button.grid(row=0, column=2, padx=6)

    # ========================================================
    # RETURN TAB
    # ========================================================

    def build_return_tab(self):

        frame = ctk.CTkFrame(
            self.return_tab
        )

        frame.pack(
            fill="both",
            expand=True,
            padx=20,
            pady=20
        )

        self.return_file_label = (
            ctk.CTkLabel(
                frame,
                text="No file selected"
            )
        )

        self.return_file_label.pack(
            pady=(20, 5)
        )

        browse = ctk.CTkButton(
            frame,
            text="Select Excel File",
            command=self.select_return_file
        )

        browse.pack(
            pady=10
        )

        self.return_validation_label = ctk.CTkLabel(
            frame, text="Select a file to validate GSTINs."
        )
        self.return_validation_label.pack(pady=(0, 4))

        # ----------------------------------------------------
        # FY
        # ----------------------------------------------------

        options_frame = ctk.CTkFrame(
            frame,
            fg_color="transparent"
        )

        options_frame.pack(
            pady=20
        )

        fy_label = ctk.CTkLabel(
            options_frame,
            text="Financial Year"
        )

        fy_label.grid(
            row=0,
            column=0,
            padx=10,
            pady=5
        )

        self.fy_combo = (
            ctk.CTkComboBox(
                options_frame,
                values=FINANCIAL_YEARS,
                width=150
            )
        )

        self.fy_combo.grid(
            row=1,
            column=0,
            padx=10
        )

        if FINANCIAL_YEARS:
            self.fy_combo.set(
                FINANCIAL_YEARS[0]
            )

        # ----------------------------------------------------
        # MONTH
        # ----------------------------------------------------

        month_label = ctk.CTkLabel(
            options_frame,
            text="Month"
        )

        month_label.grid(
            row=0,
            column=1,
            padx=10,
            pady=5
        )

        self.month_combo = (
            ctk.CTkComboBox(
                options_frame,
                values=MONTHS,
                width=150
            )
        )

        self.month_combo.grid(
            row=1,
            column=1,
            padx=10
        )

        self.month_combo.set(
            "August"
        )

        # ----------------------------------------------------
        # RETURN TYPE
        # ----------------------------------------------------

        return_label = ctk.CTkLabel(
            options_frame,
            text="Return"
        )

        return_label.grid(
            row=0,
            column=2,
            padx=10,
            pady=5
        )

        self.return_combo = (
            ctk.CTkComboBox(
                options_frame,
                values=[
                    "Both",
                    "GSTR-1",
                    "GSTR-3B"
                ],
                width=150
            )
        )

        self.return_combo.grid(
            row=1,
            column=2,
            padx=10
        )

        self.return_combo.set(
            "Both"
        )

        # ----------------------------------------------------
        # PROGRESS
        # ----------------------------------------------------

        self.return_progress = (
            ctk.CTkProgressBar(
                frame,
                width=500
            )
        )

        self.return_progress.set(
            0
        )

        self.return_progress.pack(
            pady=(25, 5)
        )

        self.return_progress_label = (
            ctk.CTkLabel(
                frame,
                text="0 / 0"
            )
        )

        self.return_progress_label.pack()

        self.return_start_button = (
            ctk.CTkButton(
                frame,
                text="CHECK FILING STATUS",
                width=220,
                height=42,
                command=self.start_returns
            )
        )

        self.return_start_button.pack(
            pady=25
        )

        return_actions = ctk.CTkFrame(frame, fg_color="transparent")
        return_actions.pack(pady=(0, 15))

        self.return_stop_button = ctk.CTkButton(
            return_actions, text="STOP / CANCEL", width=160,
            fg_color="#B3261E", hover_color="#8C1D18",
            state="disabled", command=self.request_stop
        )
        self.return_stop_button.grid(row=0, column=0, padx=6)

        self.return_open_file_button = ctk.CTkButton(
            return_actions, text="Open Output File", width=160,
            state="disabled", command=self.open_output_file
        )
        self.return_open_file_button.grid(row=0, column=1, padx=6)

        self.return_open_folder_button = ctk.CTkButton(
            return_actions, text="Open Output Folder", width=170,
            state="disabled", command=self.open_output_folder
        )
        self.return_open_folder_button.grid(row=0, column=2, padx=6)

    # ========================================================
    # FILE SELECTION
    # ========================================================

    def _validate_selected_file(self, path, target):
        try:
            gstins = read_gstins(path)
            valid, invalid, duplicates = validate_gstins(gstins)
            if not valid:
                raise ValueError("No structurally valid GSTINs were found.")

            summary = (
                f"{len(valid)} valid  •  {len(duplicates)} duplicate(s) removed  •  "
                f"{len(invalid)} invalid"
            )
            if target == "taxpayer":
                self.taxpayer_valid_gstins = valid
                self.taxpayer_validation_label.configure(text=summary)
            else:
                self.return_valid_gstins = valid
                self.return_validation_label.configure(text=summary)

            if invalid:
                preview = "\n".join(invalid[:10])
                extra = "" if len(invalid) <= 10 else f"\n...and {len(invalid)-10} more"
                messagebox.showwarning(
                    "Invalid GSTINs",
                    f"{len(invalid)} invalid GSTIN(s) will be skipped:\n\n{preview}{extra}"
                )
            return True
        except Exception as exc:
            self.show_friendly_error(exc, "Excel validation")
            return False

    def select_taxpayer_file(self):
        path = filedialog.askopenfilename(
            title="Select GSTIN Excel File", filetypes=[("Excel Files", "*.xlsx")]
        )
        if path and self._validate_selected_file(path, "taxpayer"):
            self.taxpayer_input_file = path
            self.taxpayer_file_label.configure(text=os.path.basename(path))

    def select_return_file(self):
        path = filedialog.askopenfilename(
            title="Select GSTIN Excel File", filetypes=[("Excel Files", "*.xlsx")]
        )
        if path and self._validate_selected_file(path, "return"):
            self.return_input_file = path
            self.return_file_label.configure(text=os.path.basename(path))

    # ========================================================
    # LOGGING
    # ========================================================

    def log(self, message):

        def update():
            self.log_box.configure(
                state="normal"
            )

            self.log_box.insert(
                "end",
                message + "\n"
            )

            self.log_box.see(
                "end"
            )

            self.log_box.configure(
                state="disabled"
            )

        self.after(
            0,
            update
        )

    # ========================================================
    # CLIENT
    # ========================================================

    def get_client(self):

        if self.client is None:

            self.log(
                "Loading CAPTCHA model "
                "and connecting to GST..."
            )

            retries = int(self.settings.get("retry_count", 4))
            self.client = GSTClient(
                MODEL_PATH, captcha_retries=retries, request_retries=retries
            )

            self.log(
                "GST connection ready."
            )

        return self.client

    # ========================================================
    # STOP / OUTPUT ACTIONS
    # ========================================================

    def request_stop(self):
        if self.processing:
            self.stop_requested = True
            self.log("Stop requested. Finishing the current GSTIN safely...")
            self.taxpayer_stop_button.configure(state="disabled")
            self.return_stop_button.configure(state="disabled")

    def open_output_file(self):
        if self.last_output_path and os.path.exists(self.last_output_path):
            os.startfile(self.last_output_path)
        else:
            messagebox.showwarning("Output File", "No output file is available yet.")

    def open_output_folder(self):
        if self.last_output_path:
            folder = os.path.dirname(self.last_output_path)
            if os.path.isdir(folder):
                os.startfile(folder)
                return
        messagebox.showwarning("Output Folder", "No output folder is available yet.")

    def _set_output_buttons(self, enabled):
        state = "normal" if enabled else "disabled"
        self.taxpayer_open_file_button.configure(state=state)
        self.taxpayer_open_folder_button.configure(state=state)
        self.return_open_file_button.configure(state=state)
        self.return_open_folder_button.configure(state=state)

    # ========================================================
    # TAXPAYER PROCESSING
    # ========================================================

    def start_taxpayer(self):

        if self.processing:
            return

        if not hasattr(
            self,
            "taxpayer_input_file"
        ):
            messagebox.showwarning(
                "Select File",
                "Please select an Excel file."
            )

            return

        self.stop_requested = False
        self.last_output_path = None
        self._set_output_buttons(False)

        thread = threading.Thread(
            target=self.process_taxpayer,
            daemon=True
        )

        thread.start()

    def process_taxpayer(self):

        self.processing = True

        self.after(
            0,
            lambda:
            self.taxpayer_start_button.configure(
                state="disabled"
            )
        )

        self.after(0, lambda: self.taxpayer_stop_button.configure(state="normal"))

        try:
            gstins = self.taxpayer_valid_gstins or read_gstins(
                self.taxpayer_input_file
            )

            total = len(gstins)

            self.log(
                f"Loaded {total} GSTIN(s)."
            )

            client = self.get_client()

            records = []

            for index, gstin in enumerate(
                gstins,
                start=1
            ):
                if self.stop_requested:
                    self.log("Processing stopped by user.")
                    break

                self.log(
                    f"[{index}/{total}] "
                    f"Processing {gstin}"
                )

                try:
                    taxpayer = (
                        client.get_taxpayer_details(
                            gstin
                        )
                    )

                    goods = (
                        client.get_goods_services(
                            gstin
                        )
                    )

                    record = (
                        client.build_taxpayer_record(
                            taxpayer,
                            goods
                        )
                    )

                    records.append(
                        record
                    )

                    self.log(
                        "    Success: "
                        + record.get(
                            "Legal Name",
                            ""
                        )
                    )

                except Exception as exc:

                    records.append({
                        "GSTIN": gstin,
                        "Processing Status":
                            "Failed",
                        "Error / Remarks":
                            str(exc)
                    })

                    self.log(
                        "    FAILED: "
                        + str(exc)
                    )

                progress = (
                    index / total
                )

                self.after(
                    0,
                    lambda p=progress:
                    self.taxpayer_progress.set(
                        p
                    )
                )

                self.after(
                    0,
                    lambda i=index, t=total:
                    self.taxpayer_progress_label.configure(
                        text=f"{i} / {t}"
                    )
                )

                time.sleep(
                    float(self.settings.get("request_delay", REQUEST_DELAY))
                )

            output_path = (
                self.make_output_path(
                    "GST_Taxpayer_Details"
                )
            )

            save_taxpayer_records(
                records,
                output_path
            )

            self.last_output_path = output_path
            self.after(0, lambda: self._set_output_buttons(True))
            if self.settings.get("auto_open_excel", False) and not self.stop_requested:
                self.after(250, self.open_output_file)

            self.log(
                "Stopped safely. Partial output saved." if self.stop_requested else "Completed."
            )

            self.log(
                "Output: "
                + output_path
            )

            self.after(
                0,
                lambda:
                messagebox.showinfo(
                    "Completed",
                    ("Processing stopped safely.\nCompleted records were saved.\n\n" if self.stop_requested else "Taxpayer details completed.\n\n")
                    + f"Saved to:\n{output_path}"
                )
            )

        except Exception as exc:

            self.log(
                "ERROR: "
                + str(exc)
            )

            self.after(
                0,
                lambda e=exc: self.show_friendly_error(e, "Processing")
            )

        finally:

            self.processing = False
            self.after(0, lambda: self.taxpayer_stop_button.configure(state="disabled"))

            self.after(
                0,
                lambda:
                self.taxpayer_start_button.configure(
                    state="normal"
                )
            )

    # ========================================================
    # RETURN PROCESSING
    # ========================================================

    def start_returns(self):

        if self.processing:
            return

        if not hasattr(
            self,
            "return_input_file"
        ):
            messagebox.showwarning(
                "Select File",
                "Please select an Excel file."
            )

            return

        self.stop_requested = False
        self.last_output_path = None
        self._set_output_buttons(False)

        thread = threading.Thread(
            target=self.process_returns,
            daemon=True
        )

        thread.start()

    def process_returns(self):

        self.processing = True

        self.after(
            0,
            lambda:
            self.return_start_button.configure(
                state="disabled"
            )
        )

        self.after(0, lambda: self.return_stop_button.configure(state="normal"))

        try:
            gstins = self.return_valid_gstins or read_gstins(
                self.return_input_file
            )

            total = len(gstins)

            fy_display = (
                self.fy_combo.get()
            )

            # 2026-27 -> 2026
            fy = fy_display.split(
                "-"
            )[0]

            month = (
                self.month_combo.get()
            )

            selected_return = (
                self.return_combo.get()
            )

            self.log(
                f"Checking {total} GSTIN(s) "
                f"for {month} {fy_display}."
            )

            client = self.get_client()

            records = []

            for index, gstin in enumerate(
                gstins,
                start=1
            ):
                if self.stop_requested:
                    self.log("Processing stopped by user.")
                    break

                self.log(
                    f"[{index}/{total}] "
                    f"Checking {gstin}"
                )

                record = {
                    "GSTIN": gstin,
                    "Legal Name": "",
                    "Financial Year":
                        fy_display,
                    "Month":
                        month,
                    "GSTR-1 Status": "",
                    "GSTR-1 Filing Date": "",
                    "GSTR-3B Status": "",
                    "GSTR-3B Filing Date": "",
                    "Processing Status":
                        "Success",
                    "Error / Remarks": ""
                }

                try:
                    # Get taxpayer details first so
                    # output also contains legal name
                    taxpayer = (
                        client.get_taxpayer_details(
                            gstin
                        )
                    )

                    record[
                        "Legal Name"
                    ] = taxpayer.get(
                        "lgnm",
                        ""
                    )

                    return_data = (
                        client.get_return_details(
                            gstin,
                            fy
                        )
                    )

                    returns = (
                        client.flatten_returns(
                            return_data
                        )
                    )

                    # ----------------------------------------
                    # FIND SELECTED MONTH
                    # ----------------------------------------

                    gstr1 = None
                    gstr3b = None

                    for item in returns:

                        item_month = str(
                            item.get(
                                "taxp",
                                ""
                            )
                        ).strip()

                        return_type = str(
                            item.get(
                                "rtntype",
                                ""
                            )
                        ).strip().upper()

                        if (
                            item_month.lower()
                            !=
                            month.lower()
                        ):
                            continue

                        if return_type in [
                            "GSTR1",
                            "GSTR-1"
                        ]:
                            gstr1 = item

                        elif return_type in [
                            "GSTR3B",
                            "GSTR-3B"
                        ]:
                            gstr3b = item

                    # ----------------------------------------
                    # GSTR-1
                    # ----------------------------------------

                    if selected_return in [
                        "Both",
                        "GSTR-1"
                    ]:

                        if gstr1:
                            record[
                                "GSTR-1 Status"
                            ] = gstr1.get(
                                "status",
                                ""
                            )

                            record[
                                "GSTR-1 Filing Date"
                            ] = gstr1.get(
                                "dof",
                                ""
                            )

                        else:
                            record[
                                "GSTR-1 Status"
                            ] = (
                                "No filing record found"
                            )

                    # ----------------------------------------
                    # GSTR-3B
                    # ----------------------------------------

                    if selected_return in [
                        "Both",
                        "GSTR-3B"
                    ]:

                        if gstr3b:
                            record[
                                "GSTR-3B Status"
                            ] = gstr3b.get(
                                "status",
                                ""
                            )

                            record[
                                "GSTR-3B Filing Date"
                            ] = gstr3b.get(
                                "dof",
                                ""
                            )

                        else:
                            record[
                                "GSTR-3B Status"
                            ] = (
                                "No filing record found"
                            )

                    records.append(
                        record
                    )

                    self.log(
                        "    Success: "
                        + record[
                            "Legal Name"
                        ]
                    )

                except Exception as exc:

                    record[
                        "Processing Status"
                    ] = "Failed"

                    record[
                        "Error / Remarks"
                    ] = str(exc)

                    records.append(
                        record
                    )

                    self.log(
                        "    FAILED: "
                        + str(exc)
                    )

                progress = (
                    index / total
                )

                self.after(
                    0,
                    lambda p=progress:
                    self.return_progress.set(
                        p
                    )
                )

                self.after(
                    0,
                    lambda i=index, t=total:
                    self.return_progress_label.configure(
                        text=f"{i} / {t}"
                    )
                )

                time.sleep(
                    float(self.settings.get("request_delay", REQUEST_DELAY))
                )

            output_path = (
                self.make_output_path(
                    "GST_Return_Status"
                )
            )

            save_return_records(
                records,
                output_path
            )

            self.last_output_path = output_path
            self.after(0, lambda: self._set_output_buttons(True))
            if self.settings.get("auto_open_excel", False) and not self.stop_requested:
                self.after(250, self.open_output_file)

            self.log(
                "Stopped safely. Partial output saved." if self.stop_requested else "Completed."
            )

            self.log(
                "Output: "
                + output_path
            )

            self.after(
                0,
                lambda:
                messagebox.showinfo(
                    "Completed",
                    (("Processing stopped safely.\nCompleted records were saved." if self.stop_requested else "Return status check completed.")
                    + "\n\n"
                    + f"Saved to:\n{output_path}")
                )
            )

        except Exception as exc:

            self.log(
                "ERROR: "
                + str(exc)
            )

            self.after(
                0,
                lambda e=exc: self.show_friendly_error(e, "Processing")
            )

        finally:

            self.processing = False
            self.after(0, lambda: self.return_stop_button.configure(state="disabled"))

            self.after(
                0,
                lambda:
                self.return_start_button.configure(
                    state="normal"
                )
            )

    # ========================================================
    # NETWORK / ERRORS / SETTINGS / HELP / ABOUT
    # ========================================================

    def _start_network_check(self):
        threading.Thread(target=self._network_check_worker, daemon=True).start()

    def _network_check_worker(self):
        internet_ok = gst_ok = False
        try:
            c = socket.create_connection(("1.1.1.1", 443), timeout=3)
            c.close()
            internet_ok = True
        except OSError:
            pass
        if internet_ok:
            try:
                c = socket.create_connection((GST_HOST, 443), timeout=4)
                c.close()
                gst_ok = True
            except OSError:
                pass

        if not internet_ok:
            label, color = "● Offline", "#D32F2F"
        elif not gst_ok:
            label, color = "● Internet OK • GST Portal unavailable", "#F9A825"
        else:
            label, color = "● Online • GST Portal reachable", "#2E7D32"

        def update():
            try:
                self.network_label.configure(text=label, text_color=color)
                self.after(20000, self._start_network_check)
            except Exception:
                pass
        self.after(0, update)

    @staticmethod
    def _friendly_error_text(exc):
        low = str(exc).lower()
        if isinstance(exc, PermissionError) or "permission denied" in low:
            return "The Excel file cannot be accessed. If it is open, close it and try again."
        if any(x in low for x in ("connectionerror", "failed to establish", "name resolution",
                                  "getaddrinfo", "network is unreachable")):
            return "Internet connection unavailable. Check your connection and try again."
        if "timeout" in low or "timed out" in low:
            return "The GST portal took too long to respond. Please try again."
        if any(x in low for x in ("429", "too many requests", "502", "503", "504")):
            return "The GST portal is temporarily busy or unavailable. Please try again shortly."
        if "captcha" in low:
            return "CAPTCHA could not be solved after the configured retries. Please try again."
        if any(x in low for x in ("excel", "workbook", ".xlsx")):
            return "The Excel file could not be read. Make sure it is a valid .xlsx file and is not locked."
        return "The operation could not be completed. Select Show Details for technical information."

    def show_friendly_error(self, exc, context="Application"):
        self.last_error_details = (
            f"GST Bulk Search v{APP_VERSION}\n"
            f"Time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
            f"Operation: {context}\n\n"
            f"{type(exc).__name__}: {exc}\n\n{traceback.format_exc()}"
        )
        win = ctk.CTkToplevel(self)
        win.title("GST Bulk Search - Error")
        win.geometry("570x285")
        win.transient(self)
        win.grab_set()
        ctk.CTkLabel(win, text="Something went wrong",
                     font=ctk.CTkFont(size=20, weight="bold")).pack(pady=(22, 8))
        ctk.CTkLabel(win, text=self._friendly_error_text(exc),
                     wraplength=510, justify="left").pack(padx=25, pady=10)
        actions = ctk.CTkFrame(win, fg_color="transparent")
        actions.pack(pady=16)
        ctk.CTkButton(actions, text="Show Details",
                      command=self.show_error_details).grid(row=0,column=0,padx=5)
        ctk.CTkButton(actions, text="Copy Details",
                      command=self.copy_error_details).grid(row=0,column=1,padx=5)
        ctk.CTkButton(actions, text="Close",
                      command=win.destroy).grid(row=0,column=2,padx=5)

    def show_error_details(self):
        win = ctk.CTkToplevel(self)
        win.title("Technical Error Details")
        win.geometry("760x480")
        box = ctk.CTkTextbox(win)
        box.pack(fill="both", expand=True, padx=15, pady=15)
        box.insert("1.0", self.last_error_details or "No technical details available.")
        box.configure(state="disabled")
        ctk.CTkButton(win, text="Copy Details",
                      command=self.copy_error_details).pack(pady=(0,15))

    def copy_error_details(self):
        self.clipboard_clear()
        self.clipboard_append(self.last_error_details or "No technical details available.")
        self.update()
        messagebox.showinfo("Copied", "Error details copied to clipboard.")

    def show_settings(self):
        win = ctk.CTkToplevel(self)
        win.title("Settings")
        win.geometry("650x485")
        win.transient(self)
        ctk.CTkLabel(win, text="Settings",
                     font=ctk.CTkFont(size=24, weight="bold")).pack(pady=(20,15))
        form = ctk.CTkFrame(win)
        form.pack(fill="both", expand=True, padx=25, pady=(0,15))

        ctk.CTkLabel(form,text="Request delay (seconds)").grid(row=0,column=0,sticky="w",padx=15,pady=12)
        delay=ctk.CTkEntry(form,width=190); delay.insert(0,str(self.settings.get("request_delay",1.0)))
        delay.grid(row=0,column=1,padx=15,pady=12)

        ctk.CTkLabel(form,text="Retry count").grid(row=1,column=0,sticky="w",padx=15,pady=12)
        retries=ctk.CTkEntry(form,width=190); retries.insert(0,str(self.settings.get("retry_count",4)))
        retries.grid(row=1,column=1,padx=15,pady=12)

        ctk.CTkLabel(form,text="Output folder").grid(row=2,column=0,sticky="w",padx=15,pady=12)
        folder=ctk.CTkEntry(form,width=300)
        folder.insert(0,self.settings.get("output_folder",DEFAULT_SETTINGS["output_folder"]))
        folder.grid(row=2,column=1,padx=15,pady=12)

        def browse():
            p=filedialog.askdirectory(title="Select Output Folder",initialdir=folder.get() or application_dir())
            if p:
                folder.delete(0,"end"); folder.insert(0,p)
        ctk.CTkButton(form,text="Browse",width=75,command=browse).grid(row=2,column=2,padx=(0,12),pady=12)

        auto=ctk.BooleanVar(value=bool(self.settings.get("auto_open_excel",False)))
        ctk.CTkCheckBox(form,text="Automatically open Excel after completion",
                        variable=auto).grid(row=3,column=0,columnspan=3,sticky="w",padx=15,pady=12)

        ctk.CTkLabel(form,text="Theme").grid(row=4,column=0,sticky="w",padx=15,pady=12)
        theme=ctk.CTkComboBox(form,values=["System","Light","Dark"],width=190)
        theme.set(self.settings.get("theme","System")); theme.grid(row=4,column=1,padx=15,pady=12)

        def save():
            try:
                d=float(delay.get()); r=int(retries.get()); f=folder.get().strip()
                if not 0.2 <= d <= 30: raise ValueError("Request delay must be between 0.2 and 30 seconds.")
                if not 1 <= r <= 10: raise ValueError("Retry count must be between 1 and 10.")
                if not f: raise ValueError("Please select an output folder.")
                os.makedirs(f,exist_ok=True)
                self.settings.update({"request_delay":d,"retry_count":r,"output_folder":f,
                                      "auto_open_excel":bool(auto.get()),"theme":theme.get()})
                with open(SETTINGS_PATH,"w",encoding="utf-8") as h: json.dump(self.settings,h,indent=2)
                ctk.set_appearance_mode(theme.get())
                self.client=None
                messagebox.showinfo("Settings","Settings saved successfully.")
                win.destroy()
            except Exception as exc:
                messagebox.showerror("Settings",str(exc))
        ctk.CTkButton(form,text="Save Settings",command=save).grid(row=5,column=0,columnspan=3,pady=20)

    def show_how_to(self):
        win=ctk.CTkToplevel(self); win.title("How to Use"); win.geometry("620x500"); win.transient(self)
        ctk.CTkLabel(win,text="How to Use GST Bulk Search",
                     font=ctk.CTkFont(size=22,weight="bold")).pack(pady=(20,10))
        instructions=(
            "1. Prepare an Excel .xlsx file.\n2. Keep a column named GSTIN (recommended).\n"
            "3. Enter one GSTIN per row.\n4. Select the file in the required tab.\n"
            "5. GSTINs are validated before processing.\n"
            "6. Invalid GSTINs are skipped and duplicates are processed once.\n"
            "7. Click Start Processing / Check Filing Status.\n\n"
            "Example:\nGSTIN\n03ABCDE1234F1Z5\n07ABCDE1234F1Z9"
        )
        ctk.CTkLabel(win,text=instructions,justify="left",anchor="w",
                     wraplength=560).pack(fill="x",padx=30,pady=15)
        ctk.CTkButton(win,text="Download Excel Template",
                      command=self.create_excel_template).pack(pady=10)
        ctk.CTkButton(win,text="Close",command=win.destroy).pack(pady=(5,15))

    def create_excel_template(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            path=filedialog.asksaveasfilename(title="Save GSTIN Excel Template",
                    defaultextension=".xlsx",initialfile="GST_Bulk_Search_Template.xlsx",
                    filetypes=[("Excel Workbook","*.xlsx")])
            if not path: return
            wb=Workbook(); ws=wb.active; ws.title="GSTIN Input"; ws["A1"]="GSTIN"
            ws["A1"].font=Font(bold=True); ws["A1"].alignment=Alignment(horizontal="center")
            ws["A1"].fill=PatternFill("solid",fgColor="D9EAF7"); ws.column_dimensions["A"].width=24
            ws.freeze_panes="A2"
            for row in range(2,102): ws.cell(row=row,column=1).number_format="@"
            wb.save(path)
            messagebox.showinfo("Template Created",f"Template saved to:\n{path}")
        except Exception as exc:
            self.show_friendly_error(exc,"Create Excel template")

    def show_about(self):
        win=ctk.CTkToplevel(self); win.title("About GST Bulk Search"); win.geometry("520x460"); win.transient(self)
        ctk.CTkLabel(win,text="GST BULK SEARCH",
                     font=ctk.CTkFont(size=28,weight="bold")).pack(pady=(30,3))
        ctk.CTkLabel(win,text=f"Version {APP_VERSION}").pack()
        ctk.CTkLabel(win,text="GST Taxpayer & Return Filing Utility\nBuilt for GST professionals.",
                     justify="center",font=ctk.CTkFont(size=15)).pack(pady=20)
        ctk.CTkLabel(win,text='"Less searching. More productive work."',
                     font=ctk.CTkFont(size=14,slant="italic")).pack(pady=8)
        ctk.CTkLabel(win,text=f"Developed by Rahul Pal\nContact: {CONTACT_NUMBER}",
                     justify="center").pack(pady=15)
        ctk.CTkButton(win,text="Buy Me a Coffee / Contact Developer",
                      command=self.copy_contact).pack(pady=10)
        ctk.CTkButton(win,text="Close",command=win.destroy).pack(pady=10)

    def copy_contact(self):
        self.clipboard_clear(); self.clipboard_append(CONTACT_NUMBER); self.update()
        messagebox.showinfo("Contact Developer",
                            f"Rahul Pal\n{CONTACT_NUMBER}\n\nContact number copied to clipboard.")

    # ========================================================
    # OUTPUT PATH
    # ========================================================

    def make_output_path(
        self,
        prefix
    ):
        output_dir = self.settings.get(
            "output_folder", os.path.join(application_dir(), "output")
        )

        os.makedirs(
            output_dir,
            exist_ok=True
        )

        timestamp = (
            datetime.now().strftime(
                "%Y%m%d_%H%M%S"
            )
        )

        return os.path.join(
            output_dir,
            f"{prefix}_{timestamp}.xlsx"
        )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    app = GSTBulkApp()

    app.mainloop()